"""Document reading: contract / brief / script files -> text, content image -> OCR'd text.

Kept as small functions with no Streamlit dependency so they are testable in
isolation. read_document never fails silently: a file that cannot be read
raises UnreadableDocument with a message that is safe to show the user, and
a file that is only partly readable returns notes saying what was missed.
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field

import pdfplumber
from docx import Document as DocxDocument

from app.llm import call_vision_json

DOCUMENT_TYPES = ["pdf", "docx", "txt"]
DOCUMENT_TYPES_LABEL = "PDF, Word (.docx) or plain text (.txt)"


class UnreadableDocument(ValueError):
    """The file could not be read. The message is written for the end user."""


@dataclass
class DocumentRead:
    text: str
    notes: list[str] = field(default_factory=list)  # what could not be read, if anything


def _read_pdf(data: bytes) -> DocumentRead:
    pages: list[str] = []
    image_only: list[int] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for number, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                pages.append(text)
                if len(text) < 40 and page.images:
                    image_only.append(number)
    except Exception as e:
        detail = f"{type(e).__name__} {e}".lower()
        if "password" in detail or "encrypt" in detail:
            raise UnreadableDocument(
                "This PDF is password-protected. Remove the password and upload it again."
            ) from e
        raise UnreadableDocument(
            "This PDF could not be opened. It may be damaged. Try saving or exporting it again, "
            "or upload a Word version."
        ) from e

    text = "\n\n".join(p for p in pages if p).strip()
    if not text or (pages and len(image_only) == len(pages)):
        raise UnreadableDocument(
            "This PDF looks like a scan with no searchable text, which HYPECHECK cannot read yet. "
            "Upload a searchable PDF or a Word version instead."
        )
    notes = []
    if image_only:
        listed = ", ".join(str(n) for n in image_only)
        notes.append(
            f"Page(s) {listed} contain images or scans with no readable text. The review covers "
            "the readable pages only."
        )
    return DocumentRead(text, notes)


def _read_docx(data: bytes) -> DocumentRead:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise UnreadableDocument(
            "This Word file could not be opened. It may be damaged or password-protected. "
            "Try saving it again, or upload a PDF."
        ) from e
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    notes = []
    if doc.inline_shapes:
        notes.append("This document contains images that could not be read.")
    return DocumentRead("\n".join(parts).strip(), notes)


def _read_txt(data: bytes) -> DocumentRead:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return DocumentRead(data.decode(encoding).strip())
        except UnicodeDecodeError:
            continue
    raise UnreadableDocument("This text file could not be decoded. Save it as UTF-8 and upload it again.")


def read_document(filename: str, data: bytes) -> DocumentRead:
    """Read a PDF, DOCX or TXT upload. Raises UnreadableDocument with a user-ready message."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        result = _read_pdf(data)
    elif lower.endswith(".docx"):
        result = _read_docx(data)
    elif lower.endswith(".txt"):
        result = _read_txt(data)
    elif lower.endswith(".doc"):
        raise UnreadableDocument(
            "Older Word (.doc) files are not supported. Save the file as .docx or PDF and upload that."
        )
    else:
        extension = lower.rsplit(".", 1)[-1] if "." in lower else "this"
        raise UnreadableDocument(
            f"{extension.upper()} files are not supported. Upload {DOCUMENT_TYPES_LABEL}."
        )
    if not result.text:
        raise UnreadableDocument("No readable text was found in this file.")
    return result


_OCR_SYSTEM_PROMPT = """You transcribe the visible text from a screenshot of a social \
media post (e.g. Instagram post, reel cover, story). Return ONLY a JSON object with \
this shape, no other prose:

{
  "caption_text": "the caption / body text of the post, verbatim",
  "hashtags": ["list", "of", "hashtags", "without the # sign"],
  "on_image_text": "any text overlaid directly on the image/video itself (titles, \
captions burned into the video, disclosure labels shown on-screen), verbatim",
  "visible_disclosure_labels": ["any disclosure-looking labels you can see, e.g. \
'Paid Partnership', 'Ad', 'Sponsored' -- exactly as shown"],
  "platform_ui_notes": "brief note on what platform UI is visible, if identifiable \
(e.g. Instagram post, Instagram Story, YouTube video thumbnail) - do not include \
UI chrome (like counts, timestamps, or their own button text) in caption_text"
}

If a field has nothing to report, use an empty string or empty list. Do not \
invent text that is not actually visible in the image."""


def ocr_content_image(image_bytes: bytes, media_type: str) -> dict:
    """Use vision LLM to transcribe an uploaded screenshot into structured text
    fields (caption, hashtags, on-image text, visible disclosure labels)."""
    result = call_vision_json(
        system=_OCR_SYSTEM_PROMPT,
        user_text="Transcribe this social media post screenshot as instructed.",
        image_bytes=image_bytes,
        media_type=media_type,
    )
    return result


def content_image_to_text(ocr_result: dict) -> str:
    """Flatten the structured OCR result into a single text blob for modules
    that just need "the content as text"."""
    parts = []
    if ocr_result.get("caption_text"):
        parts.append(ocr_result["caption_text"])
    if ocr_result.get("on_image_text"):
        parts.append(f"[On-screen text]: {ocr_result['on_image_text']}")
    if ocr_result.get("hashtags"):
        parts.append(" ".join(f"#{h}" for h in ocr_result["hashtags"]))
    if ocr_result.get("visible_disclosure_labels"):
        parts.append(
            "[Visible disclosure labels]: " + ", ".join(ocr_result["visible_disclosure_labels"])
        )
    return "\n".join(parts).strip()
